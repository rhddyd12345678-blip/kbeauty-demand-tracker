"""사용자 엑셀 → 긴 형식 관측값. 파일을 raw/에 넣고 한 번 실행하면 된다(여러 번 실행해도 중복 없이 합쳐짐).

  raw/아마존_*.xlsx       'low data' 시트(date, rank, asin, brand, 한국생산, 한국브랜드, title)
                           → data/obs/amazon.csv (카테고리 '스킨케어', 출처 'Wayback 스냅샷(엑셀)')
                             + data/amazon_titles.csv 에 제품명·브랜드(엑셀에서 검토된 브랜드명) 보충
  raw/*.xlsx (네이버 데이터랩 검색어트렌드 다운로드)
                           파일 이름과 상관없이 내용(앞부분 '기간' 행 + '날짜' 머리글)으로 알아본다.
                           → data/obs/naver_excel.csv (일간은 주간 평균으로). 월 1회 갱신 방법은 README.
사용: python scripts/import_excel.py
"""
from __future__ import annotations

import csv
import sys
from collections import defaultdict
from datetime import date, timedelta
from pathlib import Path

import openpyxl

sys.path.insert(0, str(Path(__file__).resolve().parent))
from collect.amazon import FIELDS, TITLES, load_titles  # noqa: E402
from lib.core import RAW, append_obs, set_status  # noqa: E402


def amazon() -> str:
    msgs = []
    for f in sorted(RAW.glob("아마존_*.xlsx")):
        wb = openpyxl.load_workbook(f, read_only=True, data_only=True)
        if "low data" not in wb.sheetnames:
            msgs.append(f"{f.name}: 'low data' 시트 없음")
            continue
        rows, titles = [], load_titles()
        it = wb["low data"].iter_rows(values_only=True)
        head = [str(h).strip() if h else "" for h in next(it)]
        ix = {h: i for i, h in enumerate(head)}
        for r in it:
            if not r or not r[ix["date"]]:
                continue
            d = str(r[ix["date"]])[:10]
            asin = str(r[ix["asin"]]).strip()
            rows.append({"기준일": d, "축": "플라이휠", "지표": "Top50 순위", "구분": f"스킨케어|{asin}",
                         "값": int(r[ix["rank"]]), "단위": "위", "출처": f"Wayback 스냅샷({f.name})"})
            t = titles.setdefault(asin, {"asin": asin, "title": "", "byline": "", "brand": "", "first_seen": d})
            t["title"] = t.get("title") or str(r[ix["title"]] or "")
            t["brand"] = t.get("brand") or str(r[ix["brand"]] or "")
            t["first_seen"] = min(t.get("first_seen") or d, d)
        with TITLES.open("w", encoding="utf-8-sig", newline="") as fh:
            w = csv.DictWriter(fh, fieldnames=FIELDS)
            w.writeheader()
            w.writerows(sorted(({k: v.get(k, "") for k in FIELDS} for v in titles.values()), key=lambda x: x["asin"]))
        n = append_obs("amazon", rows)
        msgs.append(f"{f.name}: {len({r['기준일'] for r in rows})}개 스냅샷, +{n}행")
    return "; ".join(msgs) or "아마존 엑셀 없음"


META_KEYS = ("url", "주제", "범위", "기간", "성별", "연령대", "기기", "조회기간")


def datalab_sheets(f: Path):
    """데이터랩 다운로드 형식의 시트만 골라 (시트이름, 조건 dict, 머리글 행, 데이터 행들)을 돌려준다.
    형식 판별은 파일 이름이 아니라 내용으로: 앞부분에 '기간' 행과 '날짜' 머리글이 있으면 데이터랩 파일."""
    wb = openpyxl.load_workbook(f, read_only=True, data_only=True)
    for ws in wb.worksheets:
        grid = list(ws.iter_rows(values_only=True))
        meta, hdr = {}, None
        for i, r in enumerate(grid[:30]):
            if not r or r[0] is None:
                continue
            k = str(r[0]).strip()
            if k == "날짜":
                hdr = i
                break
            if k in META_KEYS and len(r) > 1:
                meta[k] = str(r[1] or "").strip()
        if hdr is not None and "기간" in meta:
            yield ws.title, meta, grid[hdr], grid[hdr + 1:]


def parse_series(header, rows) -> dict[str, dict[str, float]]:
    """두 가지 배치를 모두 읽는다: [날짜, 값, 날짜, 값, …] 쌍 / [날짜, 키워드1, 키워드2, …]."""
    out: dict[str, dict[str, float]] = defaultdict(dict)
    pairs = sum(1 for h in header if h == "날짜") > 1
    cols = [(c, c + 1) for c in range(0, len(header) - 1, 2)] if pairs else [(0, c) for c in range(1, len(header))]
    for r in rows:
        for dc, vc in cols:
            if vc >= len(header) or not header[vc] or dc >= len(r) or vc >= len(r):
                continue
            d, v = r[dc], r[vc]
            if not d or v in (None, ""):
                continue
            try:
                ds = str(d)[:10].replace(".", "-")
                out[str(header[vc]).strip()][ds + "-01" if len(ds) == 7 else ds] = float(v)
            except ValueError:
                continue
    return out


def naver() -> str:
    """raw/*.xlsx 중 네이버 데이터랩 다운로드 파일 → data/obs/naver_excel.csv

    한 시트 = 데이터랩 조회 1번. 값은 그 조회 안의 최댓값=100이라 같은 시트 안에서만 비교할 수 있다.
    저장: 구분 '키워드|조건', 출처 '파일#시트 (기간)'. 같은 조건·키워드로 새로 받은 파일이 들어오면
    기존 행은 남겨 두고, 사이트는 기간 끝이 가장 늦은 파일을 쓴다.
    일간 자료는 주간(월요일 시작) 평균으로 줄이고, 주간·월간은 그대로.
    """
    msgs = []
    for f in sorted(RAW.glob("*.xlsx")):
        if f.name.startswith(("아마존_", "brand_master")):
            continue
        rows, n_sheet = [], 0
        for sheet, meta, header, data in datalab_sheets(f):
            n_sheet += 1
            period = meta.get("기간", "")
            unit = "일간" if period.startswith("일간") else "주간" if period.startswith("주간") else "월간" if period.startswith("월간") else "?"
            cond = " · ".join(f"{k} {meta[k]}" for k in ("연령대", "성별", "범위", "기기") if meta.get(k))
            src = f"네이버 데이터랩 다운로드({f.name}#{sheet}, {period})"
            for kw, pts in parse_series(header, data).items():
                if unit == "일간":
                    wk: dict[str, list[float]] = defaultdict(list)
                    for d, v in pts.items():
                        day = date.fromisoformat(d)
                        wk[(day - timedelta(days=day.weekday())).isoformat()].append(v)
                    pts = {d: round(sum(v) / len(v), 3) for d, v in wk.items()}
                for d, v in pts.items():
                    rows.append({"기준일": d, "축": "개별지표", "지표": "네이버 검색지수", "구분": f"{kw}|{cond}",
                                 "값": v, "단위": f"조회 내 최댓값=100 ({'일간→주간 평균' if unit == '일간' else unit})",
                                 "출처": src})
        if n_sheet:
            n = append_obs("naver_excel", rows)
            msgs.append(f"{f.name}: 조회 {n_sheet}개, +{n}행")
    return "; ".join(msgs) or "네이버 데이터랩 엑셀 없음"


if __name__ == "__main__":
    for name, fn in (("import_amazon", amazon), ("import_naver", naver)):
        note = fn()
        set_status(name, True, note)
        print(name, note)
