"""raw/ 폴더의 엑셀을 data/*.json 긴 형식으로 변환.

지원 파일 (파일명 패턴으로 자동 인식):
  - 관세청_*.xlsx        : 관세청 수출입무역통계 '수출입 실적(품목별)' 내보내기 (월별 또는 연간)
  - 네이버트렌드_*.xlsx   : 네이버 데이터랩 트렌드 다운로드 (시트 = 연령대, 열 = 키워드)
  - 아마존_*.xlsx        : 아마존 US Top50 침투분석 (low data 시트)

사용: python scripts/ingest_excel.py
"""
from __future__ import annotations

import re
from pathlib import Path

import openpyxl

from common import RAW, load_json, save_json, merge_long, to_number, update_meta


# ---------------------------------------------------------------- 관세청
def parse_customs(path: Path) -> list[dict]:
    ws = openpyxl.load_workbook(path, read_only=True, data_only=True).worksheets[0]
    rows: list[dict] = []
    header_seen = False
    for row in ws.iter_rows(values_only=True):
        if not row or row[0] is None:
            continue
        if str(row[0]).strip() == "기간":
            header_seen = True
            continue
        if not header_seen or str(row[0]).strip() == "총계":
            continue
        period = str(row[0]).strip()  # '2022.01' 또는 '2022'
        if not re.match(r"^\d{4}(\.\d{2})?$", period):
            continue
        hs = str(row[1]).strip()
        rows.append({
            "period": period.replace(".", "-"),
            "freq": "M" if "." in period else "Y",
            "hs": hs,
            "item": str(row[2]).strip(),
            "export_ton": to_number(row[3]),
            "export_usd_k": to_number(row[4]),
            "import_ton": to_number(row[5]),
            "import_usd_k": to_number(row[6]),
            "balance_usd_k": to_number(row[7]),
            "source": path.name,
        })
    return rows


# ---------------------------------------------------------------- 네이버 트렌드
def parse_naver(path: Path) -> list[dict]:
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    rows: list[dict] = []
    for ws in wb.worksheets:
        age = ws.title.strip()
        header = None
        for row in ws.iter_rows(values_only=True):
            if row and row[0] == "날짜":
                header = list(row)
                continue
            if header is None or not row or row[0] is None:
                continue
            # 열 구조: 날짜, kw1, 날짜, kw2, ...
            for i in range(0, len(header) - 1, 2):
                kw = header[i + 1]
                d, v = row[i], row[i + 1]
                if kw is None or d is None or v is None:
                    continue
                d = str(d)[:10]
                rows.append({"date": d, "age": age, "keyword": str(kw).strip(),
                             "value": to_number(v), "source": path.name})
    return rows


# ---------------------------------------------------------------- 아마존
def parse_amazon(path: Path) -> list[dict]:
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    if "low data" not in wb.sheetnames:
        return []
    ws = wb["low data"]
    rows: list[dict] = []
    header = None
    for row in ws.iter_rows(values_only=True):
        if header is None:
            header = [str(c) for c in row]
            continue
        if not row or row[0] is None:
            continue
        rec = dict(zip(header, row))
        rows.append({
            "date": str(rec.get("date"))[:10],
            "rank": to_number(rec.get("rank")),
            "asin": rec.get("asin"),
            "brand": rec.get("brand"),
            "made_in_korea": int(to_number(rec.get("한국생산")) or 0),
            "korean_brand": int(to_number(rec.get("한국브랜드")) or 0),
            "title": (rec.get("title") or "")[:120],
            "source": path.name,
        })
    return rows


def merge_naver_columnar(existing: dict, rows: list[dict]) -> dict:
    """긴 형식 → 열 형식 {"dates":[...], "series":{"연령|키워드":[...]}} 병합 (용량 절감)."""
    table: dict[str, dict[str, float]] = {}
    for key, vals in (existing.get("series") or {}).items():
        table[key] = dict(zip(existing["dates"], vals))
    for r in rows:
        table.setdefault(f"{r['age']}|{r['keyword']}", {})[r["date"]] = r["value"]
    dates = sorted({d for s in table.values() for d in s})
    series = {k: [round(s[d], 3) if d in s else None for d in dates] for k, s in sorted(table.items())}
    return {"dates": dates, "series": series}


def main() -> None:
    customs, naver, amazon = [], [], []
    for p in sorted(RAW.glob("*.xlsx")):
        n = p.name
        if n.startswith("관세청"):
            r = parse_customs(p); customs += r; print(f"{n}: 관세청 {len(r)}행")
        elif n.startswith("네이버트렌드"):
            r = parse_naver(p); naver += r; print(f"{n}: 네이버 {len(r)}행")
        elif n.startswith("아마존"):
            r = parse_amazon(p); amazon += r; print(f"{n}: 아마존 {len(r)}행")

    if customs:
        old = load_json("trade.json", [])
        # API(data.go.kr)로 받은 월은 수동 엑셀로 덮어쓰지 않음
        api_keys = {(r["hs"], r["freq"], r["period"]) for r in old if r.get("source") == "data.go.kr"}
        customs = [r for r in customs if (r["hs"], r["freq"], r["period"]) not in api_keys]
        save_json("trade.json", merge_long(old, customs, ("hs", "freq", "period")))
        update_meta("trade", "ok", "raw 엑셀 ingest")
    if naver:
        save_json("naver_trend.json", merge_naver_columnar(load_json("naver_trend.json", {}), naver))
        update_meta("naver_trend", "ok", "raw 엑셀 ingest")
    if amazon:
        save_json("amazon_top50.json", merge_long(load_json("amazon_top50.json", []), amazon, ("date", "rank")))
        update_meta("amazon", "ok", "raw 엑셀 ingest")


if __name__ == "__main__":
    main()
